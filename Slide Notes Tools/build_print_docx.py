"""Build the printed exam notes as editable Word documents (.docx), for Word or LibreOffice.

usage: python build_print_docx.py [--only CH] [--fit | --font PT --cols N] [--outdir DIR]
       python build_print_docx.py --guides [--font PT]     (all quick guides in one reading copy)

For each chapter it builds two files (see PRINTED_NOTES.md):
  Printed Exam Notes/Ch NN Printed Notes.docx  how-to box + contents, quick guide, every slide, lookup tables
  Printed Exam Notes/Ch NN Index.docx          clue index, glossary, term index
Content comes from build_print_notes.py's functions, print_notes_data.py and quick_guides.py. The font is set once,
in the Normal style. Contents page numbers and "see page N" pointers are PAGEREF fields to bookmarks: LibreOffice
recalculates them; Word shows the stored values (from LibreOffice's layout) until you select all and press F9.
With --fit (the default) one font size is used for every chapter built: the largest at which each chapter's main
file plus index file fit print_notes_data.SHEETS_PER_CHAPTER double-sided sheets, as LibreOffice lays them out.
LibreOffice runs headless with SAL_USE_VCLPLUGIN=osx on the Mac; without it, it can't see the Mac's fonts. Elsewhere
(Linux) it uses `soffice` from the PATH, and Liberation Sans Narrow stands in for Arial Narrow (same letter widths).
"""
import glob
import os
import re
import subprocess
import sys
import tempfile

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

import build_print_notes as P
import print_notes_data as D
from quick_guides import QUICK_GUIDE

MAC = sys.platform == "darwin"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice" if MAC else "soffice"
PROFILE = os.path.join(tempfile.gettempdir(), "bio2_libreoffice_profile")   # a private LibreOffice profile
PAGE_W, MARGIN_LR, GAP = 8.5, 0.4, 0.2   # inches; the HL-2240 can't print within 0.25 in of the edge
BULLETS = 90                             # numbering id of the bullet list
VALIDATOR = glob.glob(os.path.expanduser("~/.claude/skills/synced/*/docx/scripts/office/validate.py"))
HOW_TO = [
    "Find the key term or the clue (a description) in the question.",
    "Look it up in this chapter's Index file: the term index (term → slide) or the clue index (description → answer).",
    "Go to that slide here; slides are in order, and every page header shows its slide codes.",
    "Check the look-alike table: the believable wrong answer is usually the look-alike from the same slide.",
    "Rule out answers with absolute words (only, all, always, exclusively, must, every): wrong every time on Unit 1.",
]
CODE_IN_TEXT = re.compile(r"\b(\d{2})\.(\d)\.(\d+)\b")
PPR_AFTER_BDR = {qn(f"w:{t}") for t in ("shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
                                         "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
                                         "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
                                         "suppressOverlap", "jc", "textDirection", "textAlignment", "outlineLvl")}


# ---------------------------------------------------------------- small OOXML helpers
def field(par, instr, cached):
    """A Word field (PAGEREF, STYLEREF, PAGE) with a stored result."""
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = OxmlElement("w:r")
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = cached
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run.append(el)
        par._p.append(run)


def bookmark(par, name, bid):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(bid))
    par._p.insert(1 if par._p.pPr is not None else 0, start)
    par._p.append(end)


def set_columns(section, n):
    """n columns; with 2+, Word's thin line between them (w:sep)."""
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(GAP * 1440)))
    cols.set(qn("w:sep"), "1" if n > 1 else "0")


def border_el(tag, sides, sz=4):
    """<tag> with thin single borders on the given sides (sz in eighths of a point)."""
    el = OxmlElement(tag)
    for side in sides:
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "2")
        e.set(qn("w:color"), "auto")
        el.append(e)
    return el


# bookmark names: letters and digits only (LibreOffice escapes "_" when exporting them as PDF destinations)
def bm_slide(code):
    ch, sec, n = code.split(".")
    return f"s{ch}{sec}{int(n):03d}"


def bm_guide(ch, part):
    return f"g{ch}{part}"


def bm_lookup(key):
    return "lk" + re.sub(r"[^A-Za-z]", "", key.title())


def narrow_blocks(t):
    """A wide table as one small (attribute, value) table per group, so it fits a column."""
    head = t["head"]
    rows = [r[0] if isinstance(r, tuple) else r for r in t["rows"]]
    if head[0] == "":   # groups are the columns (e.g. Moss, Fern, ...)
        return [(head[j], [[r[0], r[j]] for r in rows]) for j in range(1, len(head))]
    return [(r[0], [[head[i], r[i]] for i in range(1, len(head))]) for r in rows]   # groups are the rows


# ---------------------------------------------------------------- document
class Builder:
    def __init__(self, chapters, fs, cols, pages, kind):
        """kind: "main", "index" or "guides". pages: bookmark name -> stored page number for PAGEREF fields."""
        self.chapters, self.fs, self.cols, self.pages, self.kind = chapters, fs, cols, pages, kind
        self.doc = Document()
        self.em = fs * 20                       # one em in twips
        self.text_w = PAGE_W - 2 * MARGIN_LR    # inches
        self.col_w = (self.text_w - GAP * (cols - 1)) / cols
        self.bid = 0
        self.font = pymupdf.Font(fontfile=P.FONT_FILE)   # to measure words for table column widths
        self._page_setup()
        self._styles()
        self._bullets()
        zoom = self.doc.settings.element.find(qn("w:zoom"))   # python-docx's template leaves out the required percent
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    def _page_setup(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(PAGE_W), Inches(11)
        s.left_margin = s.right_margin = Inches(MARGIN_LR)
        s.top_margin, s.bottom_margin = Inches(0.55), Inches(0.5)
        s.header_distance, s.footer_distance = Inches(0.3), Inches(0.25)

    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = P.FONT_NAME
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), P.FONT_NAME)
        normal.font.size = Pt(self.fs)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = 1.0

        def para(name, bold=False, italic=False, scale=1.0, before=0.0, after=0.0, keep=False, hang=0.0, box=False):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
            s.font.bold = bold or None
            s.font.italic = italic or None
            if scale != 1.0:
                s.font.size = Pt(round(self.fs * scale * 4) / 4)
            pf = s.paragraph_format
            pf.space_before, pf.space_after = Pt(self.fs * before), Pt(self.fs * after)
            pf.keep_with_next = keep or None
            if hang:
                pf.left_indent, pf.first_line_indent = Twips(int(hang * self.em)), Twips(-int(hang * self.em))
            if box:   # consecutive paragraphs with the same border form one box
                ppr = s.element.get_or_add_pPr()
                bdr = border_el("w:pBdr", ("top", "left", "bottom", "right"))
                later = [c for c in ppr if c.tag in PPR_AFTER_BDR]   # the schema puts pBdr before these
                later[0].addprevious(bdr) if later else ppr.append(bdr)
            return s

        para("Notes title", bold=True, scale=1.25, after=0.2)
        para("How to", box=True)
        para("How to step", hang=1.0, box=True)
        para("Map section", bold=True, before=0.35, keep=True)
        para("Map entry", hang=0.75)
        para("Chapter heading", bold=True, scale=1.35, before=0.5, after=0.25, keep=True)
        para("Section heading", bold=True, scale=1.15, before=0.45, after=0.1, keep=True)
        para("Pointer", italic=True, after=0.1)
        para("Slide heading", before=0.3, keep=True)
        para("Bullet")
        para("Figure", hang=0.75)
        para("Correction", hang=0.75)
        para("Table text", after=0.12)
        para("Lookup heading", bold=True, scale=1.35, before=0.4, after=0.2, keep=True)
        para("Table title", before=0.45, keep=True)
        para("Group heading", bold=True, scale=1.03, before=0.4, after=0.1, keep=True)
        para("Entry", hang=0.75)
        para("Guide heading", bold=True, scale=1.25, before=0.2, after=0.15, keep=True)
        para("Guide part", bold=True, scale=1.1, before=0.4, after=0.1, keep=True)
        para("Guide subheading", bold=True, before=0.25, keep=True)
        para("Guide text", after=0.25)
        para("Guide bullet")
        st.add_style("Slide code", WD_STYLE_TYPE.CHARACTER)   # the running header reads this style:
        # slide codes and section names use it, so the header names what's on each page. It sets no bold:
        # bold in a character style flips the paragraph's bold, so slide codes get bold directly instead.
        header = st["Header"].paragraph_format.tab_stops   # the template's centre/right tabs would win
        header.clear_all()
        header.add_tab_stop(Inches(self.text_w), WD_TAB_ALIGNMENT.RIGHT)

    def _bullets(self):
        """One bullet list: • at the first level, – below, 0.8 em more indent per level."""
        numbering = self.doc.part.numbering_part.element
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(BULLETS))
        for lvl in range(9):
            l = OxmlElement("w:lvl")
            l.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•" if lvl == 0 else "–"),
                             ("w:lvlJc", "left")):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                l.append(e)
            ppr, ind = OxmlElement("w:pPr"), OxmlElement("w:ind")
            ind.set(qn("w:left"), str(int((0.75 + 0.8 * lvl) * self.em)))
            ind.set(qn("w:hanging"), str(int(0.75 * self.em)))
            ppr.append(ind)
            l.append(ppr)
            rpr, fonts = OxmlElement("w:rPr"), OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(a), P.FONT_NAME)
            rpr.append(fonts)
            l.append(rpr)
            abstract.append(l)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:   # abstractNum elements must come before num elements
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        num, ref = OxmlElement("w:num"), OxmlElement("w:abstractNumId")
        num.set(qn("w:numId"), str(BULLETS))
        ref.set(qn("w:val"), str(BULLETS))
        num.append(ref)
        numbering.append(num)

    # -------------------------------------------------------- pieces
    def para(self, style, *runs):
        """runs: (text, fmt); fmt has "b" bold, "i" italic, "code" the running-header marker style."""
        p = self.doc.add_paragraph(style=style)
        for text, fmt in runs:
            if "code" in fmt:
                r = p.add_run(text, style="Slide code")
                r.bold = "b" in fmt.replace("code", "") or None
            else:
                r = p.add_run(text)
                r.bold = "b" in fmt or None
                r.italic = "i" in fmt or None
        return p

    def mark(self, p, name):
        self.bid += 1
        bookmark(p, name, self.bid)

    def ref(self, p, name):
        """A PAGEREF field showing the page of bookmark `name`."""
        field(p, f"PAGEREF {name} \\h", self.pages.get(name, "0"))

    def bullet(self, level, *runs, style="Bullet"):
        p = self.para(style, *runs)
        num = p._p.get_or_add_pPr().get_or_add_numPr()
        num.get_or_add_ilvl().val = level
        num.get_or_add_numId().val = BULLETS
        return p

    def section(self, cols):
        self.doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(self.doc.sections[-1], cols)
        return self.doc.sections[-1]

    def table(self, head, rows, width):
        """A table with thin borders on every cell; a bold heading row (if any) repeats; rows never split."""
        body = ([head] if head else []) + rows
        n = len(body[0])
        word = lambda w, bold: self.font.text_length(w, fontsize=self.fs) / 72 * (1.1 if bold else 1.0)
        pad = 0.75 * self.fs / 72 + 0.03
        mins, amounts = [], []
        for i in range(n):
            longest = max([word(w, bool(head) and ri == 0) for ri, r in enumerate(body) for w in r[i].split()] + [0])
            mins.append(longest + pad)
            amounts.append(max(1.0, sum(len(r[i]) for r in body) / len(body)))
        spare = width - sum(mins)
        widths = ([m + spare * a / sum(amounts) for m, a in zip(mins, amounts)] if spare > 0
                  else [m * width / sum(mins) for m in mins])
        t = self.doc.add_table(rows=len(body), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False   # fixed layout
        for ci, w in enumerate(widths):
            t.columns[ci].width = Inches(w)
        tblPr = t._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), str(int(width * 1440)))
        tblW.set(qn("w:type"), "dxa")
        borders = border_el("w:tblBorders", ("top", "left", "bottom", "right", "insideH", "insideV"))
        mar = OxmlElement("w:tblCellMar")
        for side, val in (("w:left", int(0.3 * self.em)), ("w:right", int(0.3 * self.em))):
            e = OxmlElement(side)
            e.set(qn("w:w"), str(val))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        # schema order inside tblPr: ... tblW, jc, tblBorders, tblLayout, tblCellMar, tblLook
        anchor = tblPr.find(qn("w:tblLayout"))
        if anchor is None:
            anchor = tblPr.find(qn("w:tblLook"))
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            tblPr.append(borders)
        look = tblPr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(mar)
        else:
            tblPr.append(mar)
        for ri, row in enumerate(body):
            trPr = t.rows[ri]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if head and ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for ci, text in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.width = Inches(widths[ci])
                p = cell.paragraphs[0]
                p.style = self.doc.styles["Table text"]
                r = p.add_run(text)
                r.bold = (bool(head) and ri == 0) or None
        self.doc.add_paragraph(style="Table text")   # a little space after the table
        return t

    def _header(self, section, left):
        section.header.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.tab_stops.add_tab_stop(Inches(self.text_w), WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run(left + "\t")
        field(hp, 'STYLEREF "Slide code"', "")
        hp.add_run(" – ")
        field(hp, 'STYLEREF "Slide code" \\l', "")
        for run in hp.runs:
            run.font.size = Pt(9)
        r.font.size = Pt(8)
        fp = section.footer.paragraphs[0]
        fp.alignment = 1
        field(fp, "PAGE", "1")

    # -------------------------------------------------------- documents
    def build(self):
        {"main": self._main, "index": self._index, "guides": self._guides}[self.kind]()
        return self.doc

    def _title(self, text):
        self.para("Notes title", (text, ""))
        self._header(self.doc.sections[0], text)

    def _main(self):
        (c,) = self.chapters
        self._title(f"BIO II printed notes: Ch {c.ch} (main)")
        self.para("How to", ("How to use these notes in the exam", "b"))
        for i, t in enumerate(HOW_TO, 1):
            self.para("How to step", (f"{i}.\t{t}", ""))
        self.section(self.cols)
        self._contents_main(c)
        self.para("Chapter heading", (f"Chapter {c.ch}: {D.CHAPTER_NAMES[c.ch]}", ""))
        self._guide(c)
        sec = None
        for n in c.slides:
            if c.deck.section[n] != sec:
                sec = c.deck.section[n]
                self.para("Section heading", (f"{c.ch}.{sec} {c.section_name(sec)}", ""))
                if sec in QUICK_GUIDE.get(c.ch, {}).get("sections", {}):
                    p = self.para("Pointer", ("Quick guide for this section: page ", ""))
                    self.ref(p, bm_guide(c.ch, sec))
            self._slide(c, n, c.section_name(sec))
        self._lookup_tables(c)

    def _contents_main(self, c):
        self.para("Map section", ("Contents", ""))
        g = QUICK_GUIDE.get(c.ch)
        if g:
            self.para("Map section", ("Quick guide", ""))
            entries = [("The big picture", bm_guide(c.ch, 0))]
            entries += [(f"{c.ch}.{s} {c.section_name(s)}", bm_guide(c.ch, s)) for s in g["sections"]]
            entries += [("Look these up, don't memorize", bm_guide(c.ch, 9))]
            for label, name in entries:
                self._entry(label, name)
        sec = None
        for grp in P.groups(c):
            if grp["sec"] != sec:
                sec = grp["sec"]
                self.para("Map section", (f"{c.ch}.{sec} {c.section_name(sec)}", ""))
            a, b = c.code(grp["first"]), c.code(grp["last"])
            rng = a if a == b else f"{a}–{b.split('.')[-1]}"
            self._entry(f"{rng} {grp['base']}", bm_slide(a))
        present = self._lookup_keys(c)
        if present:
            self.para("Map section", ("Lookup tables", ""))
            for key in present:
                self._entry(dict(P.LOOKUP)[key], bm_lookup(key))
        self.para("Map section", (f"Indexes: see the Ch {c.ch} Index file (clue index, glossary, term index)", ""))

    def _entry(self, label, name):
        p = self.para("Map entry", (f"{label}\t", ""))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(self.col_w), WD_TAB_ALIGNMENT.RIGHT)
        self.ref(p, name)

    def _guide(self, c):
        g = QUICK_GUIDE.get(c.ch)
        if not g:
            return
        p = self.para("Guide heading", ("Quick guide: the big picture", "code"))
        self.mark(p, bm_guide(c.ch, 0))
        for text in g["intro"]:
            self.para("Guide text", (text, ""))
        for sec, items in g["sections"].items():
            ns = [n for n in c.slides if c.deck.section[n] == sec]
            span = f" (slides {ns[0]}–{ns[-1]})" if ns else ""
            p = self.para("Guide part", (f"{c.ch}.{sec} {c.section_name(sec)}", "code"), (span, ""))
            self.mark(p, bm_guide(c.ch, sec))
            self._guide_items(items)
        p = self.para("Guide part", ("Look these up, don't memorize", ""))
        self.mark(p, bm_guide(c.ch, 9))
        for t in g["look_up"]:
            self.bullet(0, (t, ""), style="Guide bullet")

    def _guide_items(self, items):
        for it in items:
            if isinstance(it, tuple):
                self.para("Guide subheading", (it[0], ""))
                for b in it[1]:
                    self.bullet(0, (b, ""), style="Guide bullet")
            else:
                self.para("Guide text", (it, ""))

    def _slide(self, c, n, sec_name):
        code, title = c.code(n), c.title(n)
        runs = [(code, "codeb")]
        if title != sec_name:
            runs.append((f" {title}", "b"))
        if c.frame(n):
            runs.append((f" · {c.frame(n)}", ""))
        p = self.para("Slide heading", *runs)
        if not (c.items(n) or c.figure(n) or (c.ch, n) in D.CORRECTIONS):
            p.paragraph_format.keep_with_next = False
        self.mark(p, bm_slide(code))
        for it in c.items(n):
            if it[0] == "p":
                self.bullet(it[1], (it[2], ""))
            elif it[0] == "cell":
                self.bullet(0, (it[1], "b"), (f": {it[2]}" if it[2] else "", ""))
            else:
                self.table(it[1], it[2], self.col_w)
        for label, text in c.figure(n):
            self.para("Figure", (f"{label}: ", "i"), (text, ""))
        if (c.ch, n) in D.CORRECTIONS:
            self.para("Correction", ("Textbook: ", "bi"), (D.CORRECTIONS[(c.ch, n)], ""))

    def _lookup_keys(self, c):
        data = {"lk-tables": D.TABLES, "lk-seq": D.SEQUENCES, "lk-ploidy": D.PLOIDY, "lk-uses": D.USES}
        return [k for k, entries in data.items() if any(t["ch"] == c.ch for t in entries)]

    def _lookup_tables(self, c):
        data = {"lk-tables": D.TABLES, "lk-seq": D.SEQUENCES, "lk-ploidy": D.PLOIDY, "lk-uses": D.USES}
        titles = dict(P.LOOKUP)
        for key in self._lookup_keys(c):
            p = self.para("Lookup heading", (titles[key], "code"))
            self.mark(p, bm_lookup(key))
            for t in data[key]:
                if t["ch"] != c.ch:
                    continue
                refs = P.fmt_refs(t["slides"]) if t.get("slides") else ""
                if t["title"] or refs:
                    self.para("Table title", (t["title"], "b"), ((" " if t["title"] else "") + refs, ""))
                if "note" in t:
                    self.para("Entry", (t["note"], ""))
                    continue
                if t.get("wide"):   # one small table per group, so it fits a column
                    for group, rows in narrow_blocks(t):
                        self.para("Group heading", (group, ""))
                        self.table(None, rows, self.col_w)
                    continue
                rows = [r[0][:-1] + [f"{r[0][-1]} ({P.fmt_refs(r[1])})"] if isinstance(r, tuple) else r
                        for r in t["rows"]]
                self.table(t["head"], rows, self.col_w)

    def _index(self):
        (c,) = self.chapters
        self._title(f"BIO II printed notes: Ch {c.ch} (index)")
        self.section(self.cols)
        titles = dict(P.LOOKUP)
        self.para("Map section", ("Contents", ""))
        for key in ("lk-clues", "lk-gloss", "lk-index"):
            self._entry(titles[key], bm_lookup(key))
        p = self.para("Lookup heading", (titles["lk-clues"], "code"))
        self.mark(p, bm_lookup("lk-clues"))
        for group, entries in D.CLUES:
            rows = [e for e in entries if e[2][0][0] == c.ch]
            if rows:
                self.para("Group heading", (group, ""))
                for q, a, s in rows:
                    self.para("Entry", (f"{q} → ", ""), (a, "b"), (f" {P.fmt_refs(s)}", ""))
        gl = P.glossary(self.chapters)
        p = self.para("Lookup heading", (titles["lk-gloss"], "code"))
        self.mark(p, bm_lookup("lk-gloss"))
        for t, d, refs, more in gl:
            self.para("Entry", (t, "b"), (f": {d} {P.fmt_refs(refs)}{more}", ""))
        p = self.para("Lookup heading", (titles["lk-index"], "code"))
        self.mark(p, bm_lookup("lk-index"))
        for t, note, where in P.term_index(self.chapters, gl):
            self.para("Entry", (t, "b"), (f"{(' ' + note) if note else ''} {where}", ""))

    def _guides(self):
        self._title("BIO II Unit 2 quick guides (reading copy)")
        self.section(self.cols)
        for c in self.chapters:
            self.para("Chapter heading", (f"Chapter {c.ch}: {D.CHAPTER_NAMES[c.ch]}", "code"))
            self._guide(c)


# ---------------------------------------------------------------- LibreOffice layout
def to_pdf(docx_path, outdir):
    """LibreOffice's layout of the document as a PDF, with bookmarks as named destinations (time-limited)."""
    # SAL_USE_VCLPLUGIN=osx: use the Mac's text engine, so headless LibreOffice sees the Mac's fonts (otherwise it
    # silently swaps in its own bundled fonts) and lays text out the way the LibreOffice window does
    fmt = 'pdf:writer_pdf_Export:{"ExportBookmarksToPDFDestination":{"type":"boolean","value":"true"}}'
    res = subprocess.run(["perl", "-e", "alarm 180; exec @ARGV", SOFFICE, "--headless", "--norestore",
                          f"-env:UserInstallation=file://{PROFILE}", "--convert-to", fmt, "--outdir", outdir,
                          docx_path], capture_output=True, text=True, env={**os.environ, **({"SAL_USE_VCLPLUGIN": "osx"} if MAC else {})})
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise SystemExit("LibreOffice didn't convert the document:\n" + res.stdout[-1000:] + res.stderr[-1000:])
    return pdf


def bookmark_pages(pdf_path):
    """Page (as text, 1-based) of every bookmark, from LibreOffice's PDF destinations."""
    with pymupdf.open(pdf_path) as pdf:
        return {name: str(d["page"] + 1) for name, d in pdf.resolve_names().items() if d.get("page", -1) >= 0}


def body_codes(page):
    """Bold slide codes in the page body (below the running header), column by column."""
    found = []
    for b in page.get_text("dict")["blocks"]:
        for line in b.get("lines", []):
            for s in line["spans"]:
                m = re.match(r"(\d{2}\.\d\.\d+)\b", s["text"].strip())   # LibreOffice joins code and title
                if m and s["flags"] & 16 and s["bbox"][1] > 0.5 * 72:
                    found.append((round(s["bbox"][0] / 150), s["bbox"][1], m.group(1)))
    return [t for _, _, t in sorted(found)]


def build_one(c, kind, size, tmp, out=None):
    """Build one file with stored page numbers that match LibreOffice's layout; returns (docx, pdf, pages)."""
    path = out or os.path.join(tmp, f"{kind}{c.ch}.docx")
    pages = {}
    for _ in range(4):   # store LibreOffice's page numbers in the fields until they stop changing
        Builder([c], size[0], size[1], pages, kind).build().save(path)
        pdf = to_pdf(path, tmp)
        new = bookmark_pages(pdf)
        if new == pages:
            break
        pages = new
    with pymupdf.open(pdf) as d:
        n = len(d)
    return path, pdf, n


def fit_size(chapters, budget, tmp):
    """Largest font (0.25 pt steps) at which every chapter's main + index file fit `budget` pages."""
    cache = {}

    def worst(q):
        if q not in cache:
            fs = q / 4
            size = (fs, 2 if fs >= 9 else 3)
            totals = []
            for c in chapters:
                n = 0
                for kind in ("main", "index"):
                    path = os.path.join(tmp, f"fit{kind}{c.ch}.docx")
                    Builder([c], fs, size[1], {}, kind).build().save(path)
                    with pymupdf.open(to_pdf(path, tmp)) as d:
                        n += len(d)
                totals.append(n)
            cache[q] = max(totals)
        return cache[q]

    lo, hi = 28, 64   # quarter points: 7 pt .. 16 pt
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if worst(mid) <= budget:
            lo = mid
        else:
            hi = mid - 1
    fs = lo / 4
    return fs, 2 if fs >= 9 else 3


# ---------------------------------------------------------------- checks
def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return P.clean(" ".join(parts))


def check_main(c, docx_path):
    """Every slide's text (raw XML), figure text, correction and quick-guide paragraph is in the main file."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    flat = docx_text(docx_path)
    problems = []
    for n in c.slides:
        fixes = D.TYPO_FIX.get((c.ch, n), [])
        for p in c.deck.prs.slides[n - 1]._element.iter(A + "p"):
            t = P.fix(P.clean(P.para_text(p)), fixes)
            if P.keep(t) and t not in flat and not P.CODE.match(t) and t.lower() not in c.title(n).lower():
                problems.append(f"{c.code(n)}: slide text missing from the document: {t[:80]}")
        for _, text in c.figure(n):
            if P.clean(text) not in flat:
                problems.append(f"{c.code(n)}: figure text missing from the document")
        if c.frame(n) and c.frame(n) not in flat:
            problems.append(f"{c.code(n)}: life-cycle step line missing")
    for text in guide_texts(c.ch):
        if P.clean(text) not in flat:
            problems.append(f"quick guide text missing: {text[:60]}")
    return problems


def guide_texts(ch):
    g = QUICK_GUIDE.get(ch)
    if not g:
        return []
    out = list(g["intro"]) + list(g["look_up"])
    for items in g["sections"].values():
        for it in items:
            out += [it[0]] + list(it[1]) if isinstance(it, tuple) else [it]
    return out


def check_guides(chapters):
    """Every slide code a quick guide cites exists, is in the right chapter, and mentions the statement's words."""
    STOP = P.STOP | set("they their there these those what when where while about after other also than then "
                        "through which would could every three slide slides plant plants called".split())
    problems = []
    for c in chapters:
        texts = {c.code(n): c.text(n).lower() for n in c.slides}
        for statement in re.split(r"(?<=[.!?])\s+(?=[A-Z])", " ".join(guide_texts(c.ch))):
            codes = [".".join(m.groups()) for m in CODE_IN_TEXT.finditer(statement)]
            if not codes:
                continue
            bad = [k for k in codes if k not in texts and k.startswith(f"{c.ch}.")]
            if bad:
                problems.append(f"Ch {c.ch} guide cites slides that don't exist: {bad} in: {statement[:70]}")
                continue
            words = {w for w in re.findall(r"[a-z]{5,}", CODE_IN_TEXT.sub("", statement).lower()) if w not in STOP}
            cited = " ".join(texts.get(k, "") for k in codes)
            if words and not any(w[:6] in cited for w in words):
                problems.append(f"Ch {c.ch} guide: no cited slide mentions this: {statement[:90]}")
    return problems


def check_layout(pdf_path, main=True):
    """LibreOffice's layout: no pictures, Letter pages, the chosen font, running headers (main file)."""
    problems = []
    with pymupdf.open(pdf_path) as pdf:
        fonts = {f[3].split("+")[-1] for p in pdf for f in p.get_fonts()}
        names = [P.FONT_NAME] + ([P.METRIC_TWIN[P.FONT_NAME][0]] if not MAC and P.FONT_NAME in P.METRIC_TWIN else [])
        if not any(n.replace(" ", "").lower() in f.replace(" ", "").lower() for n in names for f in fonts):
            problems.append(f"the text isn't in {P.FONT_NAME} (fonts used: {sorted(fonts)})")
        text = " ".join(p.get_text() for p in pdf)
        if "Reference source not found" in text or "Error: Reference" in text:
            problems.append("a page number shows a reference error")
        for i, p in enumerate(pdf, 1):
            if p.get_images():
                problems.append(f"page {i}: has a picture")
            if (round(p.rect.width), round(p.rect.height)) != (612, 792):
                problems.append(f"page {i}: not Letter size")
            if main:
                codes = body_codes(p)
                head = " ".join(b[4] for b in p.get_text("blocks") if b[3] < 0.5 * 72)
                shown = re.findall(r"\d{2}\.\d\.\d+", head)   # guide pages show section names instead
                if codes and shown and codes[0] not in shown:
                    problems.append(f"page {i}: running header doesn't show its first slide {codes[0]}")
    return problems


def check_stored_pages(docx_path, pdf_path):
    """The page numbers stored in the fields (what Word shows) match LibreOffice's layout."""
    actual = bookmark_pages(pdf_path)
    xml = Document(docx_path).element.xml
    stored = re.findall(r"PAGEREF (\w+) \\h.*?<w:t>([^<]*)</w:t>", xml, re.S)
    return [f"stored page for {name} is {pg}, LibreOffice has {actual.get(name)}" for name, pg in stored
            if actual.get(name) != pg]


def validate(path):
    if not VALIDATOR:
        return [f"(Word-format validator not found; checked only that LibreOffice converts {os.path.basename(path)})"]
    res = subprocess.run([sys.executable, VALIDATOR[0], path], capture_output=True, text=True,
                         cwd=os.path.dirname(os.path.dirname(os.path.dirname(VALIDATOR[0]))))
    return [] if "All validations PASSED" in res.stdout else [f"Word-format validation failed for {path}:\n"
                                                             + res.stdout[-1500:] + res.stderr[-500:]]


def main():
    arg = lambda flag, default=None: sys.argv[sys.argv.index(flag) + 1] if flag in sys.argv else default
    outdir = os.path.abspath(arg("--outdir", P.OUT_DIR))
    os.makedirs(outdir, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        if "--guides" in sys.argv:   # one reading copy of every quick guide
            chapters = [P.Chapter(ch) for ch in sorted(QUICK_GUIDE)]
            size = (float(arg("--font", 11)), int(arg("--cols", 2)))
            out = os.path.join(outdir, "Unit 2 Quick Guides.docx")
            Builder(chapters, size[0], size[1], {}, "guides").build().save(out)
            pdf = to_pdf(out, tmp)
            with pymupdf.open(pdf) as d:
                n = len(d)
            problems = check_guides(chapters) + check_layout(pdf, main=False) + validate(out)
            print(f"{size[0]:g} pt: {n} pages in LibreOffice -> {out}")
            print("\n".join(problems) if problems else "checks passed: guide citations match their slides; valid Word file")
            return
        only = int(arg("--only")) if "--only" in sys.argv else None
        chapters = [P.Chapter(ch) for ch in D.CHAPTERS if only is None or ch == only]
        budget = 2 * D.SHEETS_PER_CHAPTER
        if "--font" in sys.argv:
            size = (float(arg("--font")), int(arg("--cols", 2)))
        else:
            size = fit_size(chapters, budget, tmp)
        problems = P.check_citations(chapters) + check_guides(chapters)
        report = []
        for c in chapters:
            total = 0
            for kind, label in (("main", "Printed Notes"), ("index", "Index")):
                out = os.path.join(outdir, f"Ch {c.ch} {label}.docx")
                lock = os.path.join(outdir, "~$" + os.path.basename(out)[2:])
                if os.path.exists(lock) and "--force" not in sys.argv:
                    print(f"note: {lock} exists (Word lock file); if the file is open in Word, close it first")
                path, pdf, n = build_one(c, kind, size, tmp, out)
                total += n
                report.append(f"  {os.path.basename(out)}: {n} pages")
                problems += (check_main(c, path) if kind == "main" else []) + check_layout(pdf, kind == "main")
                problems += check_stored_pages(path, pdf) + validate(path)
            report.append(f"  Ch {c.ch} total: {total} pages ({(total + 1) // 2} sheets double-sided; budget {budget})")
            if total > budget:
                problems.append(f"Ch {c.ch}: {total} pages; the budget is {budget}")
    print(f"{size[0]:g} pt, {size[1]} columns:")
    print("\n".join(report))
    print("\n".join(problems) if problems else
          "checks passed: all slide text, figure text, corrections and guide text present; citations match their "
          "slides; page numbers match LibreOffice's layout; no pictures; valid Word files")


if __name__ == "__main__":
    main()
