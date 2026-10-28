"""Build study speaker notes for the chapter decks from the ch*_notes_spec.py files in this folder.

usage: python build_slide_notes.py [--dump notes.txt] [--check] [--only CH] [--out PATH] [--force]

  --only CH   write only chapter CH's deck (the others are still read, for slide codes)
  --out PATH  save that one deck to PATH instead of Speaker Notes/ (e.g. Speaker Notes/Working copy/...)
  --force     write even if a target deck is open in PowerPoint (a ~$ lock file next to it)

Every deck that has a spec, or that holds a home slide some spec points to, is rebuilt from the
untouched original in "Chapter PPTs and Other Study Materials/" and saved (replacing the previous build) as
"Chapter PPTs and Other Study Materials/Speaker Notes/<deck name> (speaker notes).pptx". The originals are
never written to. Rules are described in AGENTS.md, section "Slide speaker notes".
"""
import glob
import importlib.util
import os
import re
import sys

from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK_DIR = os.path.join(HERE, "..", "Chapter PPTs and Other Study Materials")
OUT_DIR = os.path.join(DECK_DIR, "Speaker Notes")
DECKS = {
    20: "Ch20_prokarya_viruses.pptx",
    21: "Ch21_Protist.pptx",
    22: "Ch22_fungi.pptx",
    23: "Mader_Biology14e_PPT_Ch23_ACCESS.pptx",
    24: "Mader_Biology14e_PPT_Ch24_ACCESS.pptx",
    25: "Mader_Biology2024Release_PPT_Ch25_ACCESS.pptx",   # 2024 Release edition, like the textbook
    26: "Mader_Biology14e_PPT_Ch26_ACCESS.pptx",
    27: "Mader_Biology14e_PPT_Ch27_ACCESS.pptx",
}
HEADERS = ["CLARIFICATION", "IMAGE DESCRIPTION", "NOT EXPLAINED ON ANY SLIDE", "NOT EXPLAINED YET",
           "HOME NOTE", "ADDS TO", "COVERED ELSEWHERE", "EXAM FOCUS"]
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
SKIP_LINES = {"Return to parent-slide containing images."}


# ---------------------------------------------------------------- loading
def load_specs():
    specs = {}
    for path in sorted(glob.glob(os.path.join(HERE, "ch*_notes_spec.py"))):
        name = os.path.splitext(os.path.basename(path))[0]
        spec = importlib.util.spec_from_file_location(name, path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        specs[mod.CHAPTER] = mod
    return specs


def output_path(ch):
    return os.path.join(OUT_DIR, os.path.splitext(DECKS[ch])[0] + " (speaker notes).pptx")


def slide_title(slide):
    if slide.shapes.title is not None and slide.shapes.title.text_frame.text.strip():
        return slide.shapes.title.text_frame.text.strip().splitlines()[0]
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().splitlines()[0]
    return ""


class Deck:
    def __init__(self, ch):
        self.ch = ch
        src = os.path.join(DECK_DIR, DECKS[ch])
        self.prs = Presentation(src)
        for s in self.prs.slides:
            if s.has_notes_slide and any(h in s.notes_slide.notes_text_frame.text for h in HEADERS[2:]):
                raise SystemExit(f"{src} has generated notes; the original deck must stay untouched")
        self.titles = {n: slide_title(s) for n, s in enumerate(self.prs.slides, 1)}
        heads = [(n, int(m.group(1))) for n, t in self.titles.items() if (m := re.match(rf"{ch}\.(\d)\b", t))]
        self.section = {}
        for n in self.titles:
            before = [sec for h, sec in heads if h <= n]
            self.section[n] = before[-1] if before else heads[0][1]

    def code(self, n):
        return f"{self.ch}.{self.section[n]}.{n}"

    def title(self, n):
        return re.sub(rf"^{self.ch}\.\d\s+", "", self.titles[n])

    def hidden(self, n):
        return self.prs.slides[n - 1]._element.get("show") == "0"


# ---------------------------------------------------------------- image descriptions
def alt_texts(slide):
    res = []
    for el in slide.shapes._spTree.iter(P_NS + "cNvPr"):
        d = (el.get("descr") or "").strip()
        if d and d not in res:
            res.append(d)
    return res


def text_alt_body(slide):
    lines = []
    for sh in slide.shapes:
        if not sh.has_text_frame or sh == slide.shapes.title:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t and t not in SKIP_LINES and not t.isdigit():
                lines.append(t)
    return lines


def text_alt_links(deck, extra):
    part2n = {s.part.partname: n for n, s in enumerate(deck.prs.slides, 1)}
    links = dict(extra)
    for n, s in enumerate(deck.prs.slides, 1):
        if deck.hidden(n):
            continue
        for rel in s.part.rels.values():
            if rel.reltype.endswith("/slide") and not rel.is_external:
                links[n] = part2n[rel.target_part.partname]
    return links


def image_lines(deck, n, links, both):
    short = alt_texts(deck.prs.slides[n - 1])
    full = text_alt_body(deck.prs.slides[links[n] - 1]) if n in links else []
    if full and n not in both:
        return full
    return short + full


# ---------------------------------------------------------------- notes
def build(specs):
    concepts = {}
    for sp in specs.values():
        for key, (name, home) in sp.CONCEPTS.items():
            if key in concepts and concepts[key] != (name, home):
                raise SystemExit(f"concept {key} defined differently in two specs: {concepts[key]} vs {(name, home)}")
            concepts[key] = (name, home)

    # home notes: concept -> gathered details, and which concepts anything points to
    gathered, referenced = {}, set()
    for ch, sp in specs.items():
        for n, key, text in sp.ADDS:
            if concepts[key][1] == (ch, n):
                raise SystemExit(f"{ch}.{n}: adds to its own home ({key})")
            gathered.setdefault(key, []).append(((ch, n), text))
            referenced.add(key)
        for d in (sp.YET, sp.COVERED):
            for n, items in d.items():
                for it in items:
                    key = it[0] if isinstance(it, tuple) else it
                    if concepts[key][1] == (ch, n):
                        raise SystemExit(f"{ch}.{n}: points to itself ({key})")
                    referenced.add(key)
        referenced.update(sp.WHY)

    # rebuild every deck with a spec, plus every deck holding a home that something points to
    targets = set(specs) | {concepts[key][1][0] for key in referenced}
    decks = {ch: Deck(ch) for ch in sorted(targets)}

    def code(home):
        return decks[home[0]].code(home[1]) if home[0] in decks else f"{home[0]}.?.{home[1]}"

    def ref(key):
        name, home = concepts[key]
        return f"{name} ({code(home)})"

    def home_note(ch, n):
        lines = []
        for key, (name, home) in concepts.items():
            if home != (ch, n) or key not in referenced:
                continue
            lines.append(("body", f"• {name}"))
            extra = []
            by_source = {}
            for (sch, sn), text in sorted(gathered.get(key, [])):
                by_source.setdefault((sch, sn), []).append(text)
            for (sch, sn), texts in by_source.items():
                d = decks[sch]
                extra.append(f"   – {d.title(sn)} ({d.code(sn)}): " + " ".join(texts))
            for sp in specs.values():
                extra += [f"   – {w}" for w in sp.WHY.get(key, [])]
            if extra:
                lines.append(("body", "   Also write here:"))
                lines += [("body", e) for e in extra]
        return lines

    out = {}
    for ch, deck in decks.items():
        sp = specs.get(ch)
        links = text_alt_links(deck, getattr(sp, "EXTRA_TEXT_ALT", {})) if sp else {}
        for n in deck.titles:
            if deck.hidden(n):
                continue
            sections = []
            if sp:
                if n in sp.CLARIFY:
                    sections.append(("CLARIFICATION", [("body", sp.CLARIFY[n])]))
                img = image_lines(deck, n, links, sp.IMG_BOTH)
                if img:
                    sections.append(("IMAGE DESCRIPTION", [("body", t) for t in img]))
                if n in sp.ANY:
                    sections.append(("NOT EXPLAINED ON ANY SLIDE", [("body", f"• {t}") for t in sp.ANY[n]]))
                if n in sp.YET:
                    sections.append(("NOT EXPLAINED YET", [("body", f"• {t} ({code(concepts[k][1])})") for k, t in sp.YET[n]]))
            hn = home_note(ch, n)
            if hn:
                sections.append(("HOME NOTE", hn))
            if sp:
                adds = []
                for sn, key, _ in sp.ADDS:
                    if sn == n and ref(key) not in adds:
                        adds.append(ref(key))
                if adds:
                    sections.append(("ADDS TO", [("body", " · ".join(adds))]))
                if n in sp.COVERED:
                    sections.append(("COVERED ELSEWHERE", [("body", " · ".join(ref(k) for k in sp.COVERED[n]))]))
                exam = getattr(sp, "EXAM", {}).get(n)
                if exam:
                    sections.append(exam_section(exam))
            if sections:
                out[(ch, n)] = sections
    return decks, out


def exam_section(exam):
    """EXAM FOCUS: (priority, one line) or (priority, likely, unlikely, how, write)."""
    priority, *lines = exam
    if priority not in ("High", "Medium", "Low"):
        raise SystemExit(f"bad exam priority {priority!r}")
    head = f"EXAM FOCUS ({priority})"
    if len(lines) == 1:
        return head, [("body", lines[0])]
    if len(lines) != 4:
        raise SystemExit(f"exam entry needs 1 or 4 lines: {exam!r}")
    out = []
    for label, text in zip(("Likely", "Unlikely", "How", "Write"), lines):
        if isinstance(text, str):
            out.append(("body", f"{label}: {text}"))
        elif label == "Write":   # first item is the table's title, the rest are its rows
            out.append(("body", f"{label}: {text[0]}"))
            out += [("body", f"   • {t}") for t in text[1:]]
        else:
            out.append(("body", f"{label}:"))
            out += [("body", f"   • {t}") for t in text]
    return head, out


def write_notes(slide, sections):
    tf = slide.notes_slide.notes_text_frame
    first = not tf.text.strip()
    if not first:
        tf.add_paragraph()
    for i, (head, lines) in enumerate(sections):
        for style, text in [("head", head)] + lines:
            if style == "head" and i and not first:
                tf.add_paragraph()
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = para.add_run()
            r.text = text
            r.font.size = Pt(12)
            r.font.bold = style == "head"


def main():
    specs = load_specs()
    decks, notes = build(specs)
    dump = []
    for (ch, n), sections in sorted(notes.items()):
        dump.append(f"\n##### {decks[ch].code(n)}  {decks[ch].title(n)}")
        for head, lines in sections:
            dump.append(head)
            dump += [t for _, t in lines]
    if "--dump" in sys.argv:
        with open(sys.argv[sys.argv.index("--dump") + 1], "w") as f:
            f.write("\n".join(dump))
    if "--check" in sys.argv:
        print(f"{len(notes)} slides would get notes; nothing written")
        return
    for (ch, n), sections in notes.items():
        write_notes(decks[ch].prs.slides[n - 1], sections)
    only = int(sys.argv[sys.argv.index("--only") + 1]) if "--only" in sys.argv else None
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else None
    if out and only is None:
        raise SystemExit("--out needs --only")
    if only is not None and only not in decks:
        raise SystemExit(f"--only {only}: no spec has notes for that deck (built: {sorted(decks)})")
    # check every target before writing any, so a refused build leaves all decks as they were
    paths = {}
    for ch in decks:
        if only is not None and ch != only:
            continue
        path = os.path.abspath(out) if out else output_path(ch)
        if os.path.abspath(path) == os.path.abspath(os.path.join(DECK_DIR, DECKS[ch])):
            raise SystemExit("refusing to overwrite the original deck")
        lock = os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))
        if os.path.exists(lock) and "--force" not in sys.argv:
            raise SystemExit(f"{path} is open in PowerPoint; close it first (or pass --force)")
        paths[ch] = path
    for ch, path in paths.items():
        os.makedirs(os.path.dirname(path), exist_ok=True)
        decks[ch].prs.save(path)
        print(f"Ch {ch}: notes on {sum(1 for c, _ in notes if c == ch)} slides -> {path}")


if __name__ == "__main__":
    main()
