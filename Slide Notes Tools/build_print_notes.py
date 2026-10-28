"""Build a text-only PDF of the Unit 2 decks for printing and bringing to the exam.

usage: python build_print_notes.py [--only CH] [--fit | --font PT --cols N] [--out PATH] [--dump text.txt] [--html PATH]

  --only CH     build just chapter CH (e.g. 23); default: every chapter in print_notes_data.CHAPTERS
  --fit         (default) use the largest font that fits print_notes_data.SHEETS_PER_CHAPTER double-sided sheets
                per chapter
  --font PT     use this font size instead (with --cols, default 2)
  --out PATH    PDF path (default: "Printed Exam Notes/<name>.pdf"; each build replaces the previous one)
  --dump FILE   also write the printed slides as plain text, for review
  --html FILE   also keep the HTML that Chrome printed

Printed notes are allowed on the exams but pictures aren't, so every figure is printed as words: the deck's
full text alternative if it has one, otherwise its short alt text plus the textbook caption where that adds
something, or the textbook's description of the process when neither explains a diagram well enough
(print_notes_data.CAPTIONS / PROCESS). The lookup pages (tables, sequences, n/2n/3n list, uses, clue index,
glossary, term index) come from print_notes_data.py. The original decks are only read.
After building it checks that every slide's text and figure text is printed, that every slide cited in the lookup
pages mentions what it's cited for, and that the PDF has no pictures or lines and a correct map.
Rules are described in AGENTS.md, "Printed exam notes".
"""
import html
import os
import re
import subprocess
import sys
import tempfile

import pymupdf

import build_slide_notes as B
import print_notes_data as D

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
# The printout's font (see "Change the printed-notes font.md"): its name as apps show it, and its font file,
# which is used for the page header and page numbers. The font must be installed on this Mac.
FONT_NAME = "Arial Narrow"
FONT_FILE = "/System/Library/Fonts/Supplemental/Arial Narrow.ttf"
# Off the Mac (e.g. a Linux cloud session): Liberation Sans Narrow has Arial Narrow's exact letter widths, and
# LibreOffice substitutes it for Arial Narrow there, so page counts match the Mac's
METRIC_TWIN = {"Arial Narrow": ("Liberation Sans Narrow",
                                "/usr/share/fonts/truetype/liberation/LiberationSansNarrow-Regular.ttf")}
if not os.path.exists(FONT_FILE) and FONT_NAME in METRIC_TWIN and os.path.exists(METRIC_TWIN[FONT_NAME][1]):
    FONT_FILE = METRIC_TWIN[FONT_NAME][1]
OUT_DIR = os.path.join(B.HERE, "..", "Printed Exam Notes")
JUNK = re.compile(r"^(copyright|access the text alternative|see separate flexart|end of main content)"
                  r"|^(figure|table) \d+\.\w+\b|^figure \d+[a-z]$|^table (summarizes|divided)|^\d+$"
                  r"|^the following content is arranged like a table"   # accessibility note before a table
                  # photo credits printed on Ch 24-27 slides
                  r"|getty images$|alamy stock photo$|/superstock$|^©|^steven p\. lynch$|university of florida$"
                  r"|^\(photo\)", re.I)
CREDIT = re.compile(r"^\([^)]*\):")             # photo credits: "(a): Steven P. Lynch; (b): ..."
CODE = re.compile(r"^\d{2}\.\d\.\d+$")
STOP = set("the a an of in and or to is are be by on at as for with which that this its it from into step".split())


# ---------------------------------------------------------------- reading the decks
def clean(text):
    return re.sub(r"\s+", " ", text.replace("\v", " ")).strip()


def keep(text):
    return bool(text) and not JUNK.search(text) and not CREDIT.match(text)


def tokens(s):
    return {w for w in re.findall(r"[a-z0-9]+", s.lower()) if w not in STOP and len(w) > 2}


def covered(sentence, text_tokens, share=0.85):
    tk = tokens(sentence)
    return not tk or len(tk & text_tokens) / len(tk) >= share


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def para_text(p):
    """A paragraph's text in order: runs, line breaks (as spaces) and equations (a14:m, e.g. "(H2C=CH2)", "H+")."""
    out = []

    def walk(el):
        if el.tag in (A + "t", M + "t"):
            out.append(el.text or "")
        elif el.tag == A + "br":
            out.append(" ")
        elif el.tag == M + "d":   # a bracketed expression
            out.append("(")
            for child in el:
                walk(child)
            out.append(")")
        elif el.tag not in (A + "rPr", A + "pPr", A + "endParaRPr"):
            for child in el:
                walk(child)

    walk(p)
    return "".join(out)


def slide_shapes(slide):
    """The slide's shapes, plus text boxes that python-pptx skips because they sit inside mc:AlternateContent
    (PowerPoint wraps a text box that contains an equation that way)."""
    from pptx.shapes.autoshape import Shape
    shapes = list(slide.shapes)
    for ac in slide.shapes._spTree.iterchildren(MC + "AlternateContent"):
        choice = ac.find(MC + "Choice")
        if choice is not None:
            shapes += [Shape(sp, slide.shapes) for sp in choice.iterchildren(PML + "sp")]
    return shapes


def sorted_shapes(shapes):
    return sorted(shapes, key=lambda sh: ((sh.top or 0), (sh.left or 0)))


def shape_items(shape, title):
    """Text of one shape as items: ("p", level, text) or ("table", rows of cells, each a list of paragraphs)."""
    if shape.shape_type == 6:   # group
        return [it for sh in sorted_shapes(shape.shapes) for it in shape_items(sh, title)]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[[clean(para_text(p._p)) for p in c.text_frame.paragraphs] for c in r.cells] for r in shape.table.rows]
        return [("table", [r for r in rows if any(any(c) for c in r)])]
    if shape.has_text_frame and shape != title:
        return [("p", p.level, clean(para_text(p._p))) for p in shape.text_frame.paragraphs
                if keep(clean(para_text(p._p)))]
    return []


def fix(text, fixes):
    for a, b in fixes:
        text = text.replace(a, b)
    return text


class Chapter:
    def __init__(self, ch):
        self.ch = ch
        self.deck = B.Deck(ch)
        self.spec = B.load_specs().get(ch)
        extra = getattr(self.spec, "EXTRA_TEXT_ALT", {}) if self.spec else {}
        self.links = B.text_alt_links(self.deck, extra)
        self.links.update(D.TEXT_ALT_MATCH.get(ch, {}))
        self.links = {n: h for n, h in self.links.items() if h}   # None: a mislabeled text alternative is unlinked
        self.both = set(getattr(self.spec, "IMG_BOTH", set()) if self.spec else set())
        self.both |= {n for c, n in D.PRINT_BOTH if c == ch}
        self.slides = [n for n in self.deck.titles if self.content(n)]

    def content(self, n):
        return not self.deck.hidden(n) and n > 2 and "end of main content" not in self.deck.titles[n].lower()

    def code(self, n):
        return self.deck.code(n)

    def title(self, n):
        """The whole title (Deck.titles keeps only its first line), without the leading section number."""
        first = clean(self.deck.titles[n])
        t = self.deck.prs.slides[n - 1].shapes.title
        full = clean(" ".join(para_text(p) for p in t._element.iter(A + "p"))) if t is not None else ""
        return re.sub(rf"^{self.ch}\.\d\s+", "", full if full.startswith(first) else first)

    def items(self, n):
        """("p", level, text), ("table", head, rows) or ("cell", heading, text) for one-column tables."""
        if (self.ch, n) in D.SLIDE_TEXT:   # text split around pictures of symbols, rebuilt (see the data file)
            return [("p", 0, t) if isinstance(t, str) else ("p",) + t for t in D.SLIDE_TEXT[(self.ch, n)]]
        slide = self.deck.prs.slides[n - 1]
        raw = [it for sh in sorted_shapes(slide_shapes(slide)) for it in shape_items(sh, slide.shapes.title)]
        title = self.title(n).lower()
        fixes = D.TYPO_FIX.get((self.ch, n), [])
        out = []
        for it in raw:
            if it[0] == "p":
                t = it[2]
                if t.lower() == title or title.startswith(t.lower()):   # a picture slide repeating its title
                    continue
                out.append(("p", it[1], fix(t, fixes)))
            elif len(it[1][0]) == 1:   # one-column table: each cell is a heading, a blank line, then its text
                for (paras,) in it[1]:
                    paras = [fix(p, fixes) for p in paras]
                    split = paras.index("") if "" in paras else len(paras)
                    head = "; ".join(p for p in paras[:split] if p)
                    out.append(("cell", head, " ".join(p for p in paras[split:] if p)))
            else:
                rows = [[fix(" ".join(p for p in c if p), fixes) for c in r] for r in it[1]]
                out.append(("table", rows[0], rows[1:]))
        levels = [it[1] for it in out if it[0] == "p"]
        low = min(levels) if levels else 0
        return [("p", it[1] - low, it[2]) if it[0] == "p" else it for it in out]

    def has_picture(self, n):
        return any(sh.shape_type == 13 for sh in self.deck.prs.slides[n - 1].shapes)

    def frame(self, n):
        """The one-line pointer for an earlier frame of a life-cycle figure, else None."""
        cycle = D.cycle_of(self.ch, n)
        if cycle and n != cycle["final"] and self.has_picture(n):
            return D.frame_line(self.ch, n, cycle, self.code(cycle["final"]))
        return None

    def figure(self, n):
        """Figure text for slide n: a list of (label, text); empty if the slide has no picture."""
        if not self.has_picture(n):
            return []
        if self.frame(n):
            return [("Fig", D.FRAME_EXTRA[(self.ch, n)])] if (self.ch, n) in D.FRAME_EXTRA else []
        fixes = D.FIG_FIX.get((self.ch, n), [])
        lines = [fix(clean(t), fixes) for t in B.image_lines(self.deck, n, self.links, self.both)]
        out = [("Fig", " ".join(t for t in lines if keep(t)))]
        if (self.ch, n) in D.CAPTIONS:
            out.append(("Textbook caption", D.CAPTIONS[(self.ch, n)]))
        if (self.ch, n) in D.PROCESS:
            out.append(("Textbook", D.PROCESS[(self.ch, n)]))
        return out

    def section_name(self, sec):
        for t in self.deck.titles.values():
            if re.match(rf"{self.ch}\.{sec}\b", t):
                return clean(re.sub(rf"^{self.ch}\.{sec}\s*", "", t))
        return ""

    def text(self, n):
        """Everything printed for slide n, as one string (for the index search and the citation check)."""
        parts = [self.title(n)]
        for it in self.items(n):
            if it[0] == "p":
                parts.append(it[2])
            elif it[0] == "cell":
                parts += [it[1], it[2]]
            else:
                parts += it[1] + [c for r in it[2] for c in r]
        parts += [t for _, t in self.figure(n)] + [self.frame(n) or ""]
        return " ".join(parts)


_DECKS = {}


def code_of(ch, n):
    if ch not in _DECKS:
        _DECKS[ch] = B.Deck(ch)
    return _DECKS[ch].code(n)


def fmt_refs(refs):
    """[(23, 17), (23, 18), (23, 21)] -> "23.2.17–18, 23.2.21" (runs of slides in one section are joined)."""
    runs = []
    for ch, n in refs:
        sec = code_of(ch, n).split(".")[1]
        if runs and runs[-1][0] == ch and runs[-1][1] == sec and n == runs[-1][3] + 1:
            runs[-1][3] = n
        else:
            runs.append([ch, sec, n, n])
    return ", ".join(code_of(ch, a) + ("" if a == b else f"–{b}") for ch, _, a, b in runs)


# ---------------------------------------------------------------- HTML
def esc(s):
    return html.escape(s, quote=False)


def anchor(code):
    return "s" + code.replace(".", "-")


# page margins leave room for the header and page number inside the HL-2240's printable area
CSS = """
@page { size: Letter; margin: 0.48in 0.36in 0.42in 0.36in; }
* { box-sizing: border-box; }
body { margin: 0; font-family: "FONTNAME", sans-serif; font-size: FSpt; line-height: 1.17; color: #000; }
.cols { column-count: COLS; column-gap: 0.2in; orphans: 2; widows: 2; }
p, div { margin: 0; }
h1 { font-size: 1.35em; margin: 0 0 0.25em; break-after: avoid; }
h1.ch { break-before: page; }
h1.first { margin-top: 0.5em; }
h1.lk { margin-top: 0.4em; }
h2 { font-size: 1.15em; margin: 0.45em 0 0.15em; break-after: avoid; }
h3 { font-size: 1.03em; margin: 0.4em 0 0.1em; break-after: avoid; }
.sl { margin-top: 0.3em; }
.hd { break-after: avoid; }
.b, .fg, .fx, .e { padding-left: 0.75em; text-indent: -0.75em; }
.fg i, .fx b { font-style: italic; }
table { border-collapse: collapse; margin: 0.1em 0 0.15em; width: 100%; }
td, th { padding: 0 0.45em 0.18em 0; vertical-align: top; text-align: left; }
th { font-weight: bold; }
.tt { font-weight: bold; margin-top: 0.45em; break-after: avoid; }
.rf { font-weight: normal; }
tr, .e, .b { break-inside: avoid; }   /* tables may continue in the next column, but rows stay whole */
.wide { margin-bottom: 0.3em; }
.me { display: flex; justify-content: space-between; gap: 0.4em; }
.me a { color: #000; text-decoration: none; }
.pg { flex: none; width: 1.6em; text-align: right; font-variant-numeric: tabular-nums; }
.ms { font-weight: bold; margin-top: 0.35em; }
.rem { margin-bottom: 0.4em; }
.rem div { padding-left: 1em; text-indent: -1em; }
.top { font-size: 1.25em; font-weight: bold; margin-bottom: 0.2em; }
"""


def bullets(items):
    out = []
    for it in items:
        if it[0] == "p":
            mark = "•" if it[1] == 0 else "–"
            out.append(f'<div class="b" style="margin-left:{it[1] * 0.8:.1f}em">{mark} {esc(it[2])}</div>')
        elif it[0] == "cell":
            out.append(f'<div class="b">• <b>{esc(it[1])}</b>{": " + esc(it[2]) if it[2] else ""}</div>')
        else:
            out.append(table_html(it[1], it[2]))
    return out


def table_html(head, rows):
    h = "".join(f"<th>{esc(c)}</th>" for c in head)
    body = "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f"<table><thead><tr>{h}</tr></thead><tbody>{body}</tbody></table>"


def slide_html(c, n, sec_name):
    code, title = c.code(n), c.title(n)
    head = f"<b>{code}</b>" + ("" if title == sec_name else f" <b>{esc(title)}</b>")
    if c.frame(n):   # an earlier frame of a life-cycle figure: one line
        head += f" · {esc(c.frame(n))}"
    out = [f'<div class="sl" id="{anchor(code)}"><div class="hd">{head}</div>']
    out += bullets(c.items(n))
    for label, text in c.figure(n):
        out.append(f'<div class="fg"><i>{esc(label)}:</i> {esc(text)}</div>')
    if (c.ch, n) in D.CORRECTIONS:
        out.append(f'<div class="fx"><b>Textbook:</b> {esc(D.CORRECTIONS[(c.ch, n)])}</div>')
    out.append("</div>")
    return "\n".join(out)


def groups(c):
    """Map headings: consecutive slides with the same base title (e.g. "Moss Life Cycle 1-8") are one heading."""
    out = []
    for n in c.slides:
        base = re.sub(r"\s+\d+$", "", c.title(n))
        if out and out[-1]["base"] == base and out[-1]["sec"] == c.deck.section[n]:
            out[-1]["last"] = n
        else:
            out.append({"base": base, "first": n, "last": n, "sec": c.deck.section[n]})
    return out


def map_html(chapters, pages):
    out = ['<div class="cols">']
    for c in chapters:
        out.append(f'<div class="ms">Chapter {c.ch}: {esc(D.CHAPTER_NAMES[c.ch])}</div>')
        sec = None
        for g in groups(c):
            if g["sec"] != sec:
                sec = g["sec"]
                out.append(f'<div class="ms">{c.ch}.{sec} {esc(c.section_name(sec))}</div>')
            a, b = c.code(g["first"]), c.code(g["last"])
            rng = a if a == b else f"{a}–{b.split('.')[-1]}"
            pg = pages.get(anchor(a), "00")
            out.append(f'<div class="me"><a href="#{anchor(a)}">{rng} {esc(g["base"])}</a>'
                       f'<span class="pg">{pg}</span></div>')
    out.append("</div>")
    return "\n".join(out)


LOOKUP = [("lk-tables", "Comparison tables"), ("lk-seq", "Sequences"),
          ("lk-ploidy", "n / 2n / 3n and where meiosis happens"), ("lk-uses", "Uses"),
          ("lk-clues", "Clue index: description → answer"), ("lk-gloss", "Glossary: terms no slide explains"),
          ("lk-index", "Term index")]


def lookup_html(chapters):
    """The lookup pages: one continuous flow in columns; wide tables span the page."""
    chs = {c.ch for c in chapters}
    blocks = []   # (wide, html)
    titles = dict(LOOKUP)

    def section(key, items):   # the first heading spans the page, above the wide tables
        blocks.append((not blocks, f'<h1 class="lk" id="{key}">{esc(titles[key])}</h1>'))
        blocks.extend(items)

    def tables(entries):
        out = []
        for t in entries:
            if t["ch"] not in chs:
                continue
            refs = f' <span class="rf">{fmt_refs(t["slides"])}</span>' if t.get("slides") else ""
            title = f'<div class="tt">{esc(t["title"])}{refs}</div>' if t["title"] or refs else ""   # untitled: the section heading names it
            if "note" in t:
                out.append((False, f'<div class="tb">{title}<div class="e">{esc(t["note"])}</div></div>'))
                continue
            rows = []
            for r in t["rows"]:
                if isinstance(r, tuple):
                    cells, slides = r
                    r = cells[:-1] + [f"{cells[-1]} ({fmt_refs(slides)})"]
                rows.append(r)
            out.append((t.get("wide", False), f'<div class="tb">{title}{table_html(t["head"], rows)}</div>'))
        return out

    section("lk-tables", tables(D.TABLES))
    section("lk-seq", tables(D.SEQUENCES))
    section("lk-ploidy", tables(D.PLOIDY))
    section("lk-uses", tables(D.USES))
    clue = []
    for group, entries in D.CLUES:
        rows = [e for e in entries if e[2][0][0] in chs]
        if rows:
            clue.append((False, f"<h3>{esc(group)}</h3>" + "".join(
                f'<div class="e">{esc(q)} → <b>{esc(a)}</b> {fmt_refs(s)}</div>' for q, a, s in rows)))
    section("lk-clues", clue)
    gl = glossary(chapters)
    section("lk-gloss", [(False, "".join(
        f'<div class="e"><b>{esc(t)}</b>: {esc(d)} {fmt_refs(refs)}{more}</div>' for t, d, refs, more in gl))])
    section("lk-index", [(False, "".join(
        f'<div class="e"><b>{esc(t)}</b>{(" " + esc(note)) if note else ""} {where}</div>'
        for t, note, where in term_index(chapters, gl)))])

    out, cols = [], []
    for wide, h in blocks:
        if wide:
            if cols:
                out.append('<div class="cols">' + "".join(cols) + "</div>")
                cols = []
            out.append(f'<div class="wide">{h}</div>')
        else:
            cols.append(h)
    if cols:
        out.append('<div class="cols">' + "".join(cols) + "</div>")
    return "\n".join(out)


def glossary(chapters):
    """Terms no printed slide explains: the spec's ANY entries, plus YET entries whose home isn't printed."""
    printed = {c.ch for c in chapters}
    seen = {}
    for c in chapters:
        sp = c.spec
        if not sp:
            continue
        for n in c.slides:
            entries = list(sp.ANY.get(n, []))
            entries += [t for k, t in sp.YET.get(n, []) if sp.CONCEPTS[k][1][0] not in printed]
            for e in entries:
                term, _, definition = e.partition(": ")
                key = D.GLOSSARY_SAME.get(term.lower(), term.lower())
                if key not in seen:
                    seen[key] = [term, definition, []]
                elif term.lower() == key:   # the entry named like the merged one supplies its wording
                    seen[key][:2] = [term, definition]
                if (c.ch, n) not in seen[key][2]:
                    seen[key][2].append((c.ch, n))
    for c in chapters:
        for term, definition, slides in D.GLOSSARY.get(c.ch, []):
            seen.setdefault(term.lower(), [term, definition, [(c.ch, n) for n in slides]])
    return sorted(((v[0], v[1], v[2][:4], " …" if len(v[2]) > 4 else "") for v in seen.values()),
                  key=lambda e: e[0].lower())


def term_index(chapters, gloss):
    """(term, note, where): main slides first, then every other printed slide that mentions it (up to 6 in all)."""
    out = []
    texts = {(c.ch, n): c.text(n) for c in chapters for n in c.slides}
    printed = {c.ch for c in chapters}
    base = lambda t: t.split(" (")[0].lower()
    names = set()
    for term, pattern, main, note in D.INDEX:
        if not any(h[0] in printed for h in main):   # each entry belongs to its main slides' chapter
            continue
        rx = re.compile(pattern, re.I)
        first = [h for h in main if h[0] in printed]
        rest = [h for h in texts if rx.search(texts[h]) and h not in first]
        refs = first + rest
        if refs:
            names.add(base(term))
            out.append((term, note, fmt_refs(sorted(refs[:6], key=lambda h: (h not in first, h)))
                        + (" …" if len(refs) > 6 else "")))
    for term, _, refs, _ in gloss:   # glossary terms: point to the glossary
        if base(term) not in names:
            out.append((term, None, "see Glossary"))
    return sorted(out, key=lambda e: e[0].lower())


def document(chapters, size, pages):
    fs, cols = size
    css = CSS.replace("FS", str(fs)).replace("COLS", str(cols)).replace("FONTNAME", FONT_NAME)
    name = label(chapters)
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{esc(name)}</title>"
           f"<style>{css}</style></head><body>"]
    out.append(f'<div class="top">BIO II {esc(name)}</div>')
    out.append('<div class="rem"><b>Ruling out answers (Unit 1 pattern):</b>' +
               "".join(f"<div>{i}. {esc(t)}</div>" for i, t in enumerate(D.REMINDERS, 1)) + "</div>")
    out.append(map_html(chapters, pages))
    for i, c in enumerate(chapters):   # the first chapter starts right after the map; later ones on a new page
        out.append(f'<h1 class="{"first" if i == 0 else "ch"}" id="ch{c.ch}">Chapter {c.ch}: '
                   f'{esc(D.CHAPTER_NAMES[c.ch])}</h1>')
        out.append('<div class="cols">')
        sec = None
        for n in c.slides:
            if c.deck.section[n] != sec:
                sec = c.deck.section[n]
                out.append(f'<h2>{c.ch}.{sec} {esc(c.section_name(sec))}</h2>')
            out.append(slide_html(c, n, c.section_name(sec)))
        out.append("</div>")
    out.append(lookup_html(chapters))
    out.append("</body></html>")
    return "\n".join(out)


def label(chapters):
    return "printed notes: " + ("Unit 2" if len(chapters) == 5 else "Ch " + ", ".join(str(c.ch) for c in chapters))


# ---------------------------------------------------------------- PDF
def render(doc_html, pdf_path, keep_html=None):
    with tempfile.TemporaryDirectory() as tmp:
        src = keep_html or os.path.join(tmp, "notes.html")
        with open(src, "w") as f:
            f.write(doc_html)
        if os.path.exists(pdf_path):
            os.remove(pdf_path)
        res = subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                              f"--user-data-dir={tmp}/profile", f"--print-to-pdf={pdf_path}",
                              "file://" + os.path.abspath(src)], capture_output=True, text=True, timeout=180)
        if not os.path.exists(pdf_path):
            raise SystemExit("Chrome didn't write the PDF:\n" + res.stderr[-2000:])


def link_pages(pdf_path, doc_html):
    """Page (1-based) of every map link's target, from the PDF's internal links (in map order)."""
    targets = re.findall(r'<div class="me"><a href="#([^"]+)"', doc_html)
    pdf = pymupdf.open(pdf_path)
    dests = [l["page"] + 1 for p in pdf for l in p.get_links() if l["kind"] in (pymupdf.LINK_GOTO, pymupdf.LINK_NAMED)]
    pdf.close()
    if len(dests) != len(targets):
        raise SystemExit(f"map has {len(targets)} links but the PDF has {len(dests)}")
    return dict(zip(targets, (str(d) for d in dests)))


def page_count(pdf_path):
    with pymupdf.open(pdf_path) as pdf:
        return len(pdf)


def bold_codes(page):
    """Slide-header codes (bold spans like 23.2.17) on a page, column by column."""
    found = []
    for b in page.get_text("dict")["blocks"]:
        for line in b.get("lines", []):
            for s in line["spans"]:
                if CODE.match(s["text"].strip()) and s["flags"] & 16:
                    found.append((round(s["bbox"][0] / 150), s["bbox"][1], s["text"].strip()))
    return [t for _, _, t in sorted(found)]


def stamp(pdf_path, out_path, name):
    """Running header (where you are) and page number on every page; a blank page makes the count even."""
    pdf = pymupdf.open(pdf_path)
    font = pymupdf.Font(fontfile=FONT_FILE)
    short = {title: re.split(r"[:(]", title)[0].replace(" and where meiosis happens", "").strip()
             for _, title in LOOKUP}
    carry, last_code = None, None
    for i, p in enumerate(pdf):
        text = p.get_text()
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        found = [short[l] for l in lines if l in short]
        codes = bold_codes(p)
        if codes:
            first = codes[0]
            starts_with_slide = lines[0] == first or lines[0].startswith("Chapter")
            a = first if starts_with_slide or not last_code else last_code
            head = f"Ch {first.split('.')[0]} · {a} – {codes[-1]}"
            if "Ruling out answers (Unit 1 pattern):" in text:
                head = "Map · " + head
            last_code = codes[-1]
        else:
            names = ([carry] if carry and not (found and lines[0] in short) else []) + found
            head = " · ".join(dict.fromkeys(names)) or "Map"
        if found:
            carry = found[-1]
        w, h = p.rect.width, p.rect.height
        p.insert_font(fontname="AN", fontfile=FONT_FILE)
        p.insert_text((27, 22), f"BIO II · {name}", fontname="AN", fontsize=8)
        p.insert_text((w - 27 - font.text_length(head, fontsize=9), 22), head, fontname="AN", fontsize=9)
        p.insert_text((w / 2 - font.text_length(str(i + 1), fontsize=9) / 2, h - 16), str(i + 1),
                      fontname="AN", fontsize=9)
    if len(pdf) % 2:
        pdf.new_page(width=pdf[0].rect.width, height=pdf[0].rect.height)
    pdf.save(out_path, garbage=3, deflate=True)
    n = len(pdf)
    pdf.close()
    return n


def fit_size(chapters, budget):
    """Largest font (0.25 pt steps) whose PDF has at most `budget` pages; 2 columns from 9 pt up, else 3."""
    with tempfile.TemporaryDirectory() as tmp:
        raw = os.path.join(tmp, "fit.pdf")
        cache = {}

        def pages(fs):
            if fs not in cache:
                size = (fs, 2 if fs >= 9 else 3)
                render(document(chapters, size, {}), raw)
                cache[fs] = page_count(raw)
            return cache[fs]

        lo, hi = 28, 64   # quarter points: 7 pt .. 16 pt
        if pages(lo / 4) > budget:
            return lo / 4, 3
        while lo < hi:   # largest size that fits
            mid = (lo + hi + 1) // 2
            if pages(mid / 4) <= budget:
                lo = mid
            else:
                hi = mid - 1
        fs = lo / 4
        return fs, 2 if fs >= 9 else 3


def build_pdf(chapters, size, out_path, keep_html=None):
    with tempfile.TemporaryDirectory() as tmp:
        raw = os.path.join(tmp, "raw.pdf")
        pages = {}
        for _ in range(4):   # fill the map's page numbers until they stop changing
            doc_html = document(chapters, size, pages)
            render(doc_html, raw, keep_html)
            new = link_pages(raw, doc_html)
            if new == pages:
                break
            pages = new
        return stamp(raw, out_path, label(chapters)), doc_html


# ---------------------------------------------------------------- checks
def check_text(chapters, doc_html):
    """Every slide's text (raw XML), figure text and every text alternative must be printed."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    plain = lambda h: clean(html.unescape(re.sub(r"<[^>]+>", " ", h)))
    flat = plain(doc_html)
    body = doc_html.split('<h1 class="lk" id="lk-')[0]
    parts = re.split(r'<div class="sl" id="(s[\d-]+)">', body)
    printed = {parts[i]: tokens(plain(parts[i + 1])) for i in range(1, len(parts), 2)}
    problems = []
    for c in chapters:
        for n in c.slides:
            slide = c.deck.prs.slides[n - 1]
            fixes = D.TYPO_FIX.get((c.ch, n), [])
            raw = [clean(para_text(p)) for p in slide._element.iter(A + "p")]
            for wrong, _ in fixes:
                if not any(wrong in t for t in raw):
                    problems.append(f"{c.code(n)}: spelling fix not applied (text not found): {wrong[:60]}")
            for t in raw:
                t = fix(t, fixes)
                if keep(t) and t not in flat and not CODE.match(t) and t.lower() not in c.title(n).lower():
                    problems.append(f"{c.code(n)}: slide text missing: {t[:90]}")
            if c.has_picture(n):
                cyc = D.cycle_of(c.ch, n)
                own = printed.get(anchor(c.code(n)), set())
                if cyc:
                    own = own | printed.get(anchor(c.code(cyc["final"])), set())
                skip = D.FIG_SKIP.get((c.ch, n), {})
                # the description the printing rule says to print (frames: their own, covered by the final one)
                for line in B.image_lines(c.deck, n, c.links, set() if cyc else c.both):
                    for s in re.split(r"(?<=[.!?])\s+", clean(line)):
                        if keep(s) and not covered(s, own) and not any(s.startswith(p) for p in skip):
                            problems.append(f"{c.code(n)}: figure text not printed: {s[:90]}")
                if not c.figure(n) and not c.frame(n):
                    problems.append(f"{c.code(n)}: picture slide without figure text")
        used = {c.links[n] for n in c.slides if n in c.links}
        for h in c.deck.titles:
            if (c.deck.hidden(h) and "text alternative" in c.deck.titles[h].lower()
                    and h not in used and h not in D.TA_SKIP.get(c.ch, set())):
                problems.append(f"Ch {c.ch}: text alternative slide {h} not used ({c.deck.titles[h][:50]})")
    return problems


def check_citations(chapters):
    """Every slide cited on the lookup pages is a content slide that mentions what it's cited for."""
    texts = {(c.ch, n): c.text(n).lower() for c in chapters for n in c.slides}
    chs = {c.ch for c in chapters}
    stem = lambda w: re.sub(r"(s|es|a|um|i|ae)$", "", w)[:6]

    def found(part, refs):
        keys = [stem(w) for w in re.findall(r"[a-z]{4,}", part)]
        return not keys or any(all(k in texts[r] for k in keys) for r in refs)

    def hit(name, refs):
        refs = [r for r in refs if r[0] in chs]
        bad = [r for r in refs if r not in texts]
        if bad:
            return f"cites slides that aren't printed content slides: {bad}"
        name = re.sub(r"\([^)]*\)", " ", name.lower())
        if " vs " in name:   # look-alike pair: each side must be on one of the cited slides
            miss = [side for side in name.split(" vs ") if not found(side, refs)]
            return f"not on the cited slides: {miss}" if miss else None
        parts = [p.strip() for p in re.split(r"[/;,]| and | or ", name) if p.strip()]
        return None if any(found(p, refs) for p in parts) else "no cited slide mentions it"

    problems = []
    for group, entries in D.CLUES:
        for q, a, refs in entries:
            if refs[0][0] in chs and (r := hit(a, refs)):
                problems.append(f"clue index {a!r}: {r}")
    for t in D.TABLES + D.USES + D.SEQUENCES + D.PLOIDY:
        if t["ch"] not in chs:
            continue
        for row in t.get("rows", []):
            if isinstance(row, tuple) and (r := hit(row[0][0], row[1])):
                problems.append(f"{t['title']} row {row[0][0]!r}: {r}")
        problems += [f"{t['title']}: cites {s}, not a printed content slide"
                     for s in t.get("slides", []) if s[0] in chs and s not in texts]
    for term, pattern, main, note in D.INDEX:
        for ref in main:
            if ref[0] in chs and (ref not in texts or not re.search(pattern, texts[ref], re.I)):
                problems.append(f"term index {term}: its main slide {ref} doesn't mention it")
    return problems


def check_pdf(path):
    """No pictures or drawn lines, fonts embedded, Letter pages, even count, the chosen font, and a correct map."""
    problems = []
    with pymupdf.open(path) as pdf:
        # Chrome silently swaps in another font if FONT_NAME is misspelled or not installed
        body = {f[3].split("+")[-1] for p in pdf for f in p.get_fonts() if "+" in f[3]}
        if not any(FONT_NAME.replace(" ", "").lower() in b.replace(" ", "").lower() for b in body):
            problems.append(f"the text isn't in {FONT_NAME} (Chrome used {sorted(body)}): "
                            f"check FONT_NAME's spelling and that the font is installed")
        for i, p in enumerate(pdf, 1):
            if p.get_images():
                problems.append(f"page {i}: has a picture")
            if p.get_drawings():
                problems.append(f"page {i}: has drawn lines or shapes")
            if (round(p.rect.width), round(p.rect.height)) != (612, 792):
                problems.append(f"page {i}: not Letter size")
            problems += [f"page {i}: font {f[3]} not embedded" for f in p.get_fonts() if f[1] in ("n/a", "")]
        if len(pdf) % 2:
            problems.append("odd page count")
        for l in pdf[0].get_links():
            code = l["nameddest"][1:].replace("-", ".")
            if code not in bold_codes(pdf[l["page"]]):
                problems.append(f"map: {code} isn't on page {l['page'] + 1}")
    return problems


def dump_text(chapters):
    out = []
    for c in chapters:
        out.append(f"\n######## Chapter {c.ch}")
        for n in c.slides:
            out.append(f"\n## {c.code(n)}  {c.title(n)}" + (f" · {c.frame(n)}" if c.frame(n) else ""))
            for it in c.items(n):
                if it[0] == "p":
                    out.append("  " * it[1] + "- " + it[2])
                elif it[0] == "cell":
                    out.append(f"- [{it[1]}] {it[2]}")
                else:
                    out += ["  | " + " | ".join(r) for r in [it[1]] + it[2]]
            for lab, t in c.figure(n):
                out.append(f"  {lab.upper()}: {t}")
            if (c.ch, n) in D.CORRECTIONS:
                out.append("  TEXTBOOK: " + D.CORRECTIONS[(c.ch, n)])
    return "\n".join(out)


def main():
    arg = lambda flag, default=None: sys.argv[sys.argv.index(flag) + 1] if flag in sys.argv else default
    only = int(arg("--only")) if "--only" in sys.argv else None
    chapters = [Chapter(ch) for ch in D.CHAPTERS if only is None or ch == only]
    if "--dump" in sys.argv:
        with open(arg("--dump"), "w") as f:
            f.write(dump_text(chapters))
    if "--font" in sys.argv:
        size = (float(arg("--font")), int(arg("--cols", 2)))
    else:
        size = fit_size(chapters, 2 * D.SHEETS_PER_CHAPTER * len(chapters))
    name = ("Unit 2" if len(chapters) == 5 else "Ch " + ", ".join(str(c.ch) for c in chapters)) + " Printed Notes"
    out = os.path.abspath(arg("--out", os.path.join(OUT_DIR, name + ".pdf")))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    pages, doc_html = build_pdf(chapters, size, out, arg("--html"))
    print(f"{size[0]:g} pt, {size[1]} columns: {pages} pages ({pages // 2} sheets double-sided) -> {out}")
    problems = check_text(chapters, doc_html) + check_citations(chapters) + check_pdf(out)
    print("\n".join(problems) if problems else
          "checks passed: all slide text, figure text and text alternatives printed; lookup citations match "
          "their slides; no pictures or lines; fonts embedded; map pages correct")


if __name__ == "__main__":
    main()
